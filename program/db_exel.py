import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, Alignment


# ==========================================================
# LOKASI
# ==========================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

DATA_DIR = os.path.join(
    BASE_DIR,
    "data"
)

EXCEL_PATH = os.path.join(
    DATA_DIR,
    "hasil_qc.xlsx"
)


# ==========================================================
# JENIS NG
# ==========================================================

JENIS_NG = [
    "burn_mark",
    "crack",
    "defect",
    "flash",
    "short_shot"
]


# ==========================================================
# NAMA SHEET
# ==========================================================

def nama_sheet_hari_ini():

    return datetime.now().strftime(
        "%Y-%m-%d"
    )


# ==========================================================
# BUAT EXCEL OTOMATIS
# ==========================================================

def siapkan_excel():

    os.makedirs(
        DATA_DIR,
        exist_ok=True
    )

    if not os.path.exists(EXCEL_PATH):

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = nama_sheet_hari_ini()

        buat_sheet_harian(sheet)

        workbook.save(EXCEL_PATH)

        print()
        print("Excel otomatis dibuat:")
        print(EXCEL_PATH)


# ==========================================================
# BUAT SHEET HARIAN
# ==========================================================

def buat_sheet_harian(sheet):

    # ------------------------------------------------------
    # JUDUL
    # ------------------------------------------------------

    sheet["A1"] = "LAPORAN QC PLASTIC"

    sheet["A1"].font = Font(
        bold=True,
        size=16
    )


    sheet["A2"] = "Tanggal"

    sheet["B2"] = sheet.title


    # ------------------------------------------------------
    # RINGKASAN
    # ------------------------------------------------------

    sheet["A4"] = "RINGKASAN"

    sheet["A4"].font = Font(
        bold=True,
        size=13
    )


    sheet["A5"] = "TOTAL NG"

    sheet["B5"] = 0


    # ------------------------------------------------------
    # JENIS NG
    # ------------------------------------------------------

    sheet["D4"] = "JUMLAH JENIS NG"

    sheet["D4"].font = Font(
        bold=True,
        size=13
    )


    sheet["D5"] = "Jenis NG"

    sheet["E5"] = "Jumlah"


    sheet["D5"].font = Font(
        bold=True
    )

    sheet["E5"].font = Font(
        bold=True
    )


    baris = 6

    for jenis in JENIS_NG:

        sheet.cell(
            row=baris,
            column=4
        ).value = jenis

        sheet.cell(
            row=baris,
            column=5
        ).value = 0

        baris += 1


    # ------------------------------------------------------
    # DATA DETAIL
    # ------------------------------------------------------

    sheet["A13"] = "DATA DETAIL"

    sheet["A13"].font = Font(
        bold=True,
        size=13
    )


    header = [
        "No",
        "Tanggal",
        "Jam",
        "Status",
        "Jenis NG",
        "Foto NG"
    ]


    for kolom, nama in enumerate(
        header,
        start=1
    ):

        cell = sheet.cell(
            row=14,
            column=kolom
        )

        cell.value = nama

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )


    # ------------------------------------------------------
    # LEBAR KOLOM
    # ------------------------------------------------------

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 15
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 20
    sheet.column_dimensions["F"].width = 30


# ==========================================================
# BUKA DATABASE HARI INI
# ==========================================================

def buka_database_hari_ini():

    siapkan_excel()

    workbook = load_workbook(
        EXCEL_PATH
    )

    nama_sheet = nama_sheet_hari_ini()


    if nama_sheet not in workbook.sheetnames:

        sheet = workbook.create_sheet(
            nama_sheet
        )

        buat_sheet_harian(
            sheet
        )

        workbook.save(
            EXCEL_PATH
        )


    return workbook


# ==========================================================
# HITUNG DATA NG
# ==========================================================

def hitung_data(sheet):

    total_ng = 0

    jumlah_ng = {

        "burn_mark": 0,

        "crack": 0,

        "defect": 0,

        "flash": 0,

        "short_shot": 0
    }


    for row in sheet.iter_rows(
        min_row=15,
        values_only=True
    ):

        if not row[0]:
            continue


        status = str(
            row[3]
        ).lower().strip()


        jenis = str(
            row[4]
        ).lower().strip()


        if status == "ng":

            total_ng += 1


            if jenis in jumlah_ng:

                jumlah_ng[jenis] += 1


    return (
        total_ng,
        jumlah_ng
    )


# ==========================================================
# UPDATE RINGKASAN
# ==========================================================

def update_ringkasan(sheet):

    (
        total_ng,
        jumlah_ng
    ) = hitung_data(sheet)


    sheet["B5"] = total_ng


    baris = 6

    for jenis in JENIS_NG:

        sheet.cell(
            row=baris,
            column=5
        ).value = jumlah_ng[jenis]

        baris += 1


# ==========================================================
# SIMPAN NG
# ==========================================================

def simpan_data(
    status,
    jenis_ng="-",
    foto="-"
):

    status = str(
        status
    ).upper().strip()


    # ======================================================
    # KITA HANYA MENYIMPAN NG
    # ======================================================

    if status != "NG":

        return


    jenis_ng = str(
        jenis_ng
    ).lower().strip()


    if jenis_ng not in JENIS_NG:

        print(
            "Jenis NG tidak dikenal:",
            jenis_ng
        )

        return


    workbook = buka_database_hari_ini()


    sheet = workbook[
        nama_sheet_hari_ini()
    ]


    # ======================================================
    # NOMOR DATA
    # ======================================================

    nomor = 1


    for row in sheet.iter_rows(
        min_row=15,
        values_only=True
    ):

        if row[0]:

            try:

                nomor = max(
                    nomor,
                    int(row[0]) + 1
                )

            except ValueError:

                pass


    # ======================================================
    # WAKTU
    # ======================================================

    sekarang = datetime.now()


    tanggal = sekarang.strftime(
        "%Y-%m-%d"
    )


    jam = sekarang.strftime(
        "%H:%M:%S"
    )


    # ======================================================
    # TAMBAH DATA NG
    # ======================================================

    sheet.append(
        [
            nomor,
            tanggal,
            jam,
            "NG",
            jenis_ng,
            ""
        ]
    )


    # ======================================================
    # BARIS DATA
    # ======================================================

    baris_data = sheet.max_row


    # ======================================================
    # MASUKKAN GAMBAR
    # ======================================================

    if (
        foto != "-"
        and os.path.exists(foto)
    ):

        try:

            gambar = Image(foto)

            gambar.width = 180
            gambar.height = 120


            sheet.add_image(
                gambar,
                f"F{baris_data}"
            )


            sheet.row_dimensions[
                baris_data
            ].height = 95


        except Exception as e:

            print(
                "Gagal menampilkan gambar:",
                e
            )


    # ======================================================
    # UPDATE RINGKASAN
    # ======================================================

    update_ringkasan(
        sheet
    )


    # ======================================================
    # SIMPAN EXCEL
    # ======================================================

    workbook.save(
        EXCEL_PATH
    )


    print()
    print(
        "Excel: NG |",
        jenis_ng
    )